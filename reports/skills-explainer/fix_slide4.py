"""Re-lay slide 4 of the user's edited copy as a two-way comparison: .txt vs a skill (which is a .md file).

python fix_slide4.py <deck.pptx>   (edits the deck in place; shape ids are the ones in the user's copy)
"""
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

C = {k: RGBColor.from_string(v) for k, v in dict(
    white="FFFFFF", mint="2BD4A8", darkmint="198568", deep="3A0561", ink="120A1B", purple="9747FF",
    cardline="5F3181").items()}
MONO = "Consolas"

path = sys.argv[1]
prs = Presentation(path)
s = prs.slides[3]
by_id = {sh.shape_id: sh for sh in s.shapes}


def place(sid, x=None, y=None, w=None, h=None):
    sh = by_id[sid]
    for attr, v in (("left", x), ("top", y), ("width", w), ("height", h)):
        if v is not None:
            setattr(sh, attr, Inches(v))


def resize_text(sid, size):
    for p in by_id[sid].text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)


def code_box(x, y, w, h, lines, size=13.5, step=17.5):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(step)
        for t, col, bold in ln:
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = MONO, Pt(size), bold, C[col]


# Title: the slide now compares two files, and the skill is the .md one.
seg = ImageFont.truetype("C:/Windows/Fonts/segoeuib.ttf", 1000)
line1, line2a, line2b = "A .txt is just plain words,", "a skill is a .md, ", "ready to use"
for t in (line1, line2a + line2b):
    assert seg.getlength(t) * 40 / 1000 < 11.2 * 72, t
tp = by_id[2].text_frame.paragraphs
tp[0].runs[0].text = line1
tp[1].runs[0].text, tp[1].runs[1].text = line2a, line2b

CW, GAP, CY, CH = 5.567, 1.0, 2.1, 4.55
cards = {  # card, file, kind, panel, preview, (icon, fact) x3
    0.6: (5, 6, 7, 8, 9, [(10, 11), (12, 13), (14, 15)]),
    0.6 + CW + GAP: (27, 28, 29, 30, 31, [(32, 33), (34, 35), (36, 37)]),
}
previews = {
    9: [[("Trading rules. Research first. Only trade with an edge of 5+ points. "
          "Never more than 20% per market.", "white", False)]],
    31: [[("---", "cardline", False)],
         [("name: ", "darkmint", True), ("seer-trade-market", "deep", False)],
         [("description: ", "darkmint", True), ("Use to trade Seer markets", "deep", False)],
         [("---", "cardline", False)],
         [("# ", "purple", True), ("Trading rules", "purple", True)],
         [("- ", "purple", True), ("Research first", "ink", False)],
         [("- ", "purple", True), ("Edge of ", "ink", False), ("**", "purple", False), ("5+ points", "ink", True),
          ("**", "purple", False)],
         [("- ", "purple", True), ("Max ", "ink", False), ("**", "purple", False), ("20%", "ink", True),
          ("**", "purple", False), (" per market", "ink", False)]],
}
for cx, (card, file_, kind, panel, preview, facts) in cards.items():
    place(card, x=cx, y=CY, w=CW, h=CH)
    place(file_, x=cx + 0.35, y=2.28, w=CW - 0.7, h=0.5)
    resize_text(file_, 26)
    place(kind, x=cx + 0.35, y=2.8, w=CW - 0.7, h=0.34)
    resize_text(kind, 16)
    place(panel, x=cx + 0.25, y=3.2, w=CW - 0.5, h=2.2)
    el = by_id[preview]._element
    el.getparent().remove(el)
    code_box(cx + 0.45, 3.33, CW - 0.9, 2.0, previews[preview])
    for j, (icon, fact) in enumerate(facts):
        fy = 5.55 + j * 0.35
        place(icon, x=cx + 0.35, y=fy + 0.07)
        place(fact, x=cx + 0.67, y=fy, w=CW - 1.0, h=0.32)
        resize_text(fact, 16)
by_id[29].text_frame.paragraphs[0].runs[0].text = "A skill, written in Markdown (.md)"

vs = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6 + CW + GAP / 2 - 0.4), Inches(CY + CH / 2 - 0.4), Inches(0.8), Inches(0.8))
vs.fill.solid()
vs.fill.fore_color.rgb = C["mint"]
vs.line.fill.background()
vs.shadow.inherit = False
tf = vs.text_frame
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "vs"
r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Segoe UI", Pt(20), True, C["ink"]

s.notes_slide.notes_text_frame.text = (
    "Both files can hold the same words. A .txt file is just words: no headings, no lists. "
    "A skill is a Markdown (.md) file called SKILL.md. Markdown adds simple marks: # for a heading, "
    "- for a list item, ** for bold, so people and AI can see the structure. On top, a skill has a label: "
    "a name and a line that says when to use it. It sits where the agent looks for skills, so the agent "
    "can find it and open it by itself.")
prs.save(path)
print("saved", path)
