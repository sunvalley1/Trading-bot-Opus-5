"""Build the five-slide deck "What is a skill?" in Seer's colours.

python build_deck.py  ->  ../What-is-a-skill.pptx

Colours: seer.pm's stylesheet and the Seer pitch deck (deep purple grounds, mint and lavender highlights,
near-black cards). Logo and sparkle stars: seer.pm/assets/body, prepared by prep_assets.py.
The last slide's speaker notes carry the full prompt from Desktop/Sample Prompt.docx.
"""
import math
import pathlib
import random

import docx
from lxml import etree
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE.parent / "What-is-a-skill.pptx"
PROMPT_DOCX = pathlib.Path(r"C:\Users\12\Desktop\Sample Prompt.docx")

HEX = dict(
    deep="3A0561", grad="6A2798", card="120A1B", cardline="5F3181", code="221333", light="F1DFFF",
    mint="2BD4A8", darkmint="198568", lav="E0B7FF", cyan="65CDE8", white="FFFFFF", soft="CFBDE6",
    purple="9747FF", ink="120A1B",
)
C = {k: RGBColor.from_string(v) for k, v in HEX.items()}
SW, SH, LM = 13.333, 7.5, 0.6
UI, UISB, MONO = "Segoe UI", "Segoe UI Semibold", "Consolas"
FONT_FILES = {(UI, False, False): "segoeui.ttf", (UI, True, False): "segoeuib.ttf", (UI, False, True): "segoeuii.ttf",
              (UISB, False, False): "seguisb.ttf", (UISB, False, True): "seguisbi.ttf",
              (MONO, False, False): "consola.ttf", (MONO, True, False): "consolab.ttf"}
_font_cache = {}


def R(text, size, color="white", font=UI, bold=False, italic=False, spc=None):
    return dict(text=text, size=size, color=color, font=font, bold=bold, italic=italic, spc=spc)


def P(*runs, align=PP_ALIGN.LEFT, before=0, line=None):
    return dict(runs=runs, align=align, before=before, line=line)


def _width_pt(text, run):
    key = (run["font"], run["bold"], run["italic"])
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype("C:/Windows/Fonts/" + FONT_FILES[key], 1000)
    return _font_cache[key].getlength(text) * run["size"] / 1000


def _check_fit(label, paras, w, h):
    """Greedy word wrap with the real font metrics; warn when a box would overflow."""
    total = 0.0
    for p in paras:
        words = [(word, r) for r in p["runs"] for word in r["text"].split(" ")]
        lines, cur = 1, 0.0
        for word, r in words:
            ww = _width_pt(word + " ", r)
            if cur and cur + ww - _width_pt(" ", r) > w * 72:
                lines, cur = lines + 1, ww
            else:
                cur += ww
        size = max(r["size"] for r in p["runs"])
        total += p["before"] + lines * (p["line"] if p["line"] else size * 1.2)
    if total > h * 72 + 1.5:
        print(f"  ! {label}: needs {total / 72:.2f} in, box is {h:.2f} in")


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, label="text"):
    paras = [p if isinstance(p, dict) and "runs" in p else P(p) for p in paras]
    _check_fit(label, paras, w, h)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p["align"]
        if p["before"]:
            para.space_before = Pt(p["before"])
        if p["line"]:
            para.line_spacing = Pt(p["line"])
        for r in p["runs"]:
            run = para.add_run()
            run.text = r["text"]
            f = run.font
            f.name, f.size, f.bold, f.italic = r["font"], Pt(r["size"]), r["bold"], r["italic"]
            f.color.rgb = C[r["color"]]
            if r["spc"]:
                run._r.get_or_add_rPr().set("spc", str(int(r["spc"] * 100)))
    return tb


def box(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = C[fill]
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = C[line]
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def label_in(shape, runs, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor, tf.word_wrap = MSO_ANCHOR.MIDDLE, False
    para = tf.paragraphs[0]
    para.alignment = align
    for r in runs:
        run = para.add_run()
        run.text = r["text"]
        f = run.font
        f.name, f.size, f.bold, f.italic = r["font"], Pt(r["size"]), r["bold"], r["italic"]
        f.color.rgb = C[r["color"]]


def pic(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(ASSETS / name), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def star(slide, color, x, y, w):
    pic(slide, f"star-{color}.png", x, y, w=w)


def ground(slide, gradient=False):
    f = slide.background.fill
    if gradient:
        f.gradient()
        f.gradient_angle = 0
        stops = f.gradient_stops
        stops[0].position, stops[1].position = 0.0, 1.0
        stops[0].color.rgb, stops[1].color.rgb = C["deep"], C["grad"]
    else:
        f.solid()
        f.fore_color.rgb = C["deep"]


def title(slide, line1, line2):
    """Two-line slide title, same place on every content slide. Lines are lists of (text, colour)."""
    text(slide, LM, 0.42, 11.2, 1.4, [P(*[R(t, 40, c, bold=True) for t, c in ln]) for ln in (line1, line2)],
         label="title")
    star(slide, "white", 12.18, 0.5, 0.4)
    star(slide, "lav", 12.68, 1.02, 0.22)


def footer(slide, n):
    text(slide, LM, 6.92, 1.0, 0.3, [R(str(n), 12, "soft")], label="page")
    pic(slide, "seer-lockup-white.png", SW - LM - 0.81, 6.86, h=0.36)


def code_lines(slide, x, y, w, lines, size, step, label):
    """lines: list of [(text, colour, bold)]; exact line spacing so callouts can line up with rows."""
    paras = [P(*[R(t, size, c, font=MONO, bold=b) for t, c, b in ln], line=step) for ln in lines]
    text(slide, x, y, w, len(lines) * step / 72 + 0.05, paras, label=label)


def cover_background(path):
    W, H = 2667, 1500
    rnd = random.Random(7)
    row = Image.new("RGB", (W, 1))
    a, b = (0x07, 0x02, 0x36), (0x1E, 0x04, 0x5C)
    for x in range(W):
        t = x / (W - 1)
        row.putpixel((x, 0), tuple(round(a[i] + (b[i] - a[i]) * t) for i in range(3)))
    img = row.resize((W, H))

    def glow(img, cx, cy, r, color, strength):
        small = Image.new("L", (W // 4, H // 4), 0)
        ImageDraw.Draw(small).ellipse([(cx - r) / 4, (cy - r) / 4, (cx + r) / 4, (cy + r) / 4], fill=int(255 * strength))
        mask = small.filter(ImageFilter.GaussianBlur(r / 4 * 0.6)).resize((W, H), Image.BICUBIC)
        return Image.composite(Image.new("RGB", (W, H), color), img, mask)

    cx, cy = 0.79 * W, 0.49 * H
    img = glow(img, cx, cy, 0.30 * W, (0x4A, 0x0C, 0xB0), 0.85)
    img = glow(img, cx, cy, 0.13 * W, (0x93, 0x06, 0xF0), 0.55)
    rays = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(rays)
    fx, fy = 0.62 * W, 0.46 * H
    for _ in range(240):
        ang = rnd.uniform(0, 2 * math.pi)
        r0, ln = rnd.uniform(0.10, 0.62) * W, rnd.uniform(0.03, 0.26) * W
        x0, y0 = fx + r0 * math.cos(ang), fy + r0 * math.sin(ang)
        x1, y1 = fx + (r0 + ln) * math.cos(ang), fy + (r0 + ln) * math.sin(ang)
        col = rnd.choices([(101, 205, 232), (84, 120, 255), (160, 100, 255), (210, 225, 255)], weights=[4, 3, 3, 1])[0]
        alpha = rnd.randint(30, 120)
        if (x0 + x1) / 2 < 0.56 * W and 0.18 * H < (y0 + y1) / 2 < 0.82 * H:   # keep the title area calm
            alpha = int(alpha * 0.3)
        d.line([(x0, y0), (x1, y1)], fill=col + (alpha,), width=rnd.choice([1, 2, 2, 3]))
    img = Image.alpha_composite(img.convert("RGBA"), rays.filter(ImageFilter.GaussianBlur(1.2))).convert("RGB")
    img.save(path, quality=90)


def seer_theme(prs):
    """Theme colours and fonts, so anything added later in PowerPoint lands in Seer colours."""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    scheme.set("name", "Seer")
    for tag, key in dict(dk1="ink", lt1="white", dk2="deep", lt2="light", accent1="purple", accent2="mint",
                         accent3="cyan", accent4="lav", accent5="grad", accent6="cardline", hlink="cyan",
                         folHlink="lav").items():
        el = scheme.find(f"a:{tag}", ns)
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, qn("a:srgbClr")).set("val", HEX[key])
    fonts = root.find(".//a:fontScheme", ns)
    fonts.set("name", "Seer")
    for kind in ("majorFont", "minorFont"):
        fonts.find(f"a:{kind}/a:latin", ns).set("typeface", UI)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    # Dark slides: map text to the light theme colour and paint the master purple.
    prs.slide_master._element.find(qn("p:clrMap")).attrib.update(dict(bg1="dk1", tx1="lt1", bg2="dk2", tx2="lt2"))
    prs.slide_master.background.fill.solid()
    prs.slide_master.background.fill.fore_color.rgb = C["deep"]
    # The default template is 4:3; stretch master and layout placeholders to the 16:9 width.
    for el in [prs.slide_master._element] + [lay._element for lay in prs.slide_layouts]:
        for xfrm in el.iter(qn("a:xfrm")):
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            if off is not None and ext is not None:
                off.set("x", str(int(int(off.get("x")) * 4 / 3)))
                ext.set("cx", str(int(int(ext.get("cx")) * 4 / 3)))


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
seer_theme(prs)
prs.core_properties.title = "What is a skill?"
prs.core_properties.subject = "Skills for AI agents, compared with .txt and .md files, and a prompt that builds a Seer trading bot"
BLANK = prs.slide_layouts[6]

# 1 - Cover -------------------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg = HERE / "cover-bg.jpg"
cover_background(bg)
s.shapes.add_picture(str(bg), 0, 0, prs.slide_width, prs.slide_height)
pic(s, "seer-lockup-white.png", LM, 0.55, h=0.55)
text(s, LM, 2.0, 7.4, 1.15, [R("What is a skill?", 66, bold=True)], label="cover title")
text(s, LM, 3.22, 7.4, 0.55, [R("HOW AI AGENTS LEARN A JOB", 28, "cyan", bold=True, spc=1)], label="cover tagline")
text(s, LM, 3.95, 6.6, 0.8, [R("A skill vs a .txt file vs a .md file, and a real prompt that builds a Seer trading bot", 20, font=UISB)],
     label="cover sub")
text(s, LM, 6.85, 4, 0.35, [R("seer.pm", 14, "soft")], label="cover url")
pic(s, "raven-on-dark.png", 8.2, 1.3, h=4.6)
star(s, "white", 7.55, 1.0, 0.36)
star(s, "lavlight", 12.45, 0.8, 0.2)
star(s, "lav", 12.5, 5.95, 0.28)
star(s, "pink", 7.9, 6.15, 0.18)
s.notes_slide.notes_text_frame.text = (
    "Today, in simple words: what a skill is, how it is different from a .txt file and a .md file, "
    "and a real prompt we use to build a Seer trading bot.")

# 2 - What a skill is ---------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
ground(s)
title(s, [("A skill is a ", "white"), ("how-to guide", "mint")], [("for an AI agent", "white")])
text(s, LM, 1.95, 11.5, 0.45, [R("Like a recipe card: one job, step by step, in a file called SKILL.md", 20, "lav")],
     label="s2 sub")
box(s, LM, 2.65, 6.9, 4.0, fill="card", line="cardline")
text(s, 1.0, 2.88, 6.2, 0.35, [P(R("seer-trade-market/", 15, "soft", font=MONO), R("SKILL.md", 15, "mint", font=MONO, bold=True))],
     label="s2 path")
STEP = 21.6
y0 = 3.42
code_lines(s, 1.0, y0, 6.3, [
    [("---", "soft", False)],
    [("name: ", "mint", False), ("seer-trade-market", "white", False)],
    [("description: ", "mint", False), ("Use whenever someone asks to", "white", False)],
    [("  trade, bet on, or price Seer markets.", "white", False)],
    [("---", "soft", False)],
    [("# Seer: trade a market", "lav", True)],
    [("1. Research the question", "white", False)],
    [("2. Estimate your own probability", "white", False)],
    [("3. Reconcile it with the market price", "white", False)],
    [("4. Trade only what survives", "white", False)],
], 15, STEP, "s2 code")
row = STEP / 72
box(s, 7.62, y0 + 0.04, 0.16, 5 * row - 0.08, line="mint", line_w=1.5, shape=MSO_SHAPE.RIGHT_BRACKET)
box(s, 7.62, y0 + 5 * row + 0.04, 0.16, 5 * row - 0.08, line="lav", line_w=1.5, shape=MSO_SHAPE.RIGHT_BRACKET)
text(s, 7.98, y0 + 0.02, 4.75, 1.5, [
    P(R("THE LABEL", 14, "mint", bold=True, spc=1.5)),
    P(R("A name, and when to use it", 24, font=UISB), before=4)], label="s2 label")
text(s, 7.98, y0 + 5 * row + 0.02, 4.75, 1.5, [
    P(R("THE GUIDE", 14, "lav", bold=True, spc=1.5)),
    P(R("The steps to do the job, in plain words", 24, font=UISB), before=4)], label="s2 guide")
text(s, 7.98, 6.12, 4.75, 0.5, [R("A skill can also carry scripts and notes.", 15, "soft")], label="s2 extras")
footer(s, 2)
s.notes_slide.notes_text_frame.text = (
    "A skill is a how-to guide for one job. It lives in a folder, and the main file is SKILL.md. "
    "The top of the file is the label: a name, and a description that says when to use it. "
    "The rest is the guide: the steps, in plain words. A skill can also carry scripts the agent can run "
    "and notes it can read. This example is the seer-trade-market skill from this repository, shortened.")

# 3 - How the agent uses it ---------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
ground(s, gradient=True)
title(s, [("The agent reads the ", "white"), ("label", "mint"), (" first,", "white")],
      [("and the ", "white"), ("guide", "lav"), (" only when it is needed", "white")])
CW, GAP, CY, CH = 3.644, 0.6, 2.15, 3.75
cols = [LM + i * (CW + GAP) for i in range(3)]
steps = [("It sees the labels", "At the start it knows only each skill's name and when to use it."),
         ("You ask for a job", "You type a request. One label matches it."),
         ("It opens the guide", "It reads the full SKILL.md and follows the steps.")]
for i, (cx, (head, body)) in enumerate(zip(cols, steps)):
    box(s, cx, CY, CW, CH, fill="light")
    badge = box(s, cx + 0.3, CY + 0.3, 0.55, 0.55, fill="deep", shape=MSO_SHAPE.OVAL)
    label_in(badge, [R(str(i + 1), 20, bold=True)])
    text(s, cx + 0.3, CY + 1.0, CW - 0.6, 0.45, [R(head, 21, "deep", bold=True)], label=f"s3 head {i}")
    text(s, cx + 0.3, CY + 1.5, CW - 0.6, 0.8, [R(body, 15, "ink")], label=f"s3 body {i}")
    if i < 2:
        box(s, cx + CW + GAP / 2 - 0.13, CY + CH / 2 - 0.2, 0.26, 0.4, fill="mint", shape=MSO_SHAPE.CHEVRON)
vy = CY + 2.45
for j, name in enumerate(["seer-trade-market", "seer-create-market", "pptx"]):
    pill = box(s, cols[0] + 0.3, vy + j * 0.36, 2.25, 0.29, fill="white", line="deep", line_w=0.75, radius=0.145)
    label_in(pill, [R(name, 12, "deep", font=MONO)])
bubble = box(s, cols[1] + 0.3, vy, CW - 0.6, 0.48, fill="deep", radius=0.2)
label_in(bubble, [R("\u201cTrade this Seer market.\u201d", 14, font=UISB)])
match = box(s, cols[1] + 0.3, vy + 0.66, 2.25, 0.32, fill="mint", radius=0.16)
label_in(match, [R("seer-trade-market", 12, "ink", font=MONO, bold=True)])
doc = box(s, cols[2] + 0.3, vy, CW - 0.6, 1.0, fill="white", line="deep", line_w=0.75, radius=0.1)
text(s, cols[2] + 0.48, vy + 0.1, CW - 0.9, 0.85, [
    P(R("SKILL.md", 12, "deep", font=MONO, bold=True), line=17),
    P(R("1. Research   2. Estimate", 12, "ink", font=MONO), line=17),
    P(R("3. Reconcile  4. Trade", 12, "ink", font=MONO), line=17)], label="s3 doc")
star(s, "mint", LM, 6.2, 0.32)
text(s, LM + 0.48, 6.13, 11.0, 0.5, [R("So one agent can hold many skills and still stay focused.", 22, font=UISB)],
     label="s3 statement")
footer(s, 3)
s.notes_slide.notes_text_frame.text = (
    "The agent does not read every skill all the time. At the start it sees only the labels: each skill's "
    "name and description. When your request matches a label, it opens that skill and follows the steps. "
    "That is why one agent can carry many skills and still stay focused on the job in front of it.")

# 4 - .txt vs .md vs skill ----------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
ground(s)
title(s, [("A .txt is plain, a .md is organized,", "white")], [("a skill is ", "white"), ("ready to use", "mint")])
CY, CH = 2.1, 4.55
kinds = [
    dict(file="notes.txt", kind="Plain text", dark=True, lines=[
        [("Trading rules. Research", "white", False)],
        [("first. Only trade with an", "white", False)],
        [("edge of 5+ points. Never", "white", False)],
        [("more than 20% per market.", "white", False)]],
        facts=["Just words, no marks", "You paste it in", "Good for quick notes"]),
    dict(file="notes.md", kind="Markdown", dark=True, lines=[
        [("# ", "lav", True), ("Trading rules", "white", True)],
        [("- ", "lav", True), ("Research first", "white", False)],
        [("- ", "lav", True), ("Edge of ", "white", False), ("**", "lav", False), ("5+ points", "white", True), ("**", "lav", False)],
        [("- ", "lav", True), ("Max ", "white", False), ("**", "lav", False), ("20%", "white", True), ("**", "lav", False), (" per market", "white", False)]],
        facts=["Words + marks: # - **", "You paste it in", "Good for docs and plans"]),
    dict(file="SKILL.md", kind="A skill", dark=False, lines=[
        [("---", "cardline", False)],
        [("name: ", "darkmint", True), ("seer-trade-market", "deep", False)],
        [("description: ", "darkmint", True), ("Use to", "deep", False)],
        [("  trade Seer markets", "deep", False)],
        [("---", "cardline", False)],
        [("# Trading rules", "purple", True)],
        [("- Research first \u2026", "ink", False)]],
        facts=["Has a \u201cuse when\u201d label","The agent opens it itself", "Good for jobs you repeat"]),
]
for i, (cx, k) in enumerate(zip(cols, kinds)):
    dark = k["dark"]
    box(s, cx, CY, CW, CH, fill="card" if dark else "light", line="cardline" if dark else None)
    text(s, cx + 0.3, CY + 0.22, CW - 0.6, 0.45, [R(k["file"], 22, "white" if dark else "deep", font=MONO, bold=True)],
         label=f"s4 file {i}")
    text(s, cx + 0.3, CY + 0.7, CW - 0.6, 0.32, [R(k["kind"], 14, "soft" if dark else "cardline")], label=f"s4 kind {i}")
    box(s, cx + 0.22, CY + 1.12, CW - 0.44, 1.92, fill="code" if dark else "white", radius=0.08)
    code_lines(s, cx + 0.38, CY + 1.25, CW - 0.7, k["lines"], 12.5, 17, f"s4 code {i}")
    for j, fact in enumerate(k["facts"]):
        fy = CY + 3.27 + j * 0.4
        star(s, "mint" if dark else "primary", cx + 0.3, fy + 0.06, 0.2)
        text(s, cx + 0.62, fy, CW - 0.85, 0.36, [R(fact, 15, "white" if dark else "ink")], label=f"s4 fact {i}.{j}")
footer(s, 4)
s.notes_slide.notes_text_frame.text = (
    "All three files can hold the same words. A .txt file is just words: no headings, no lists. "
    "A .md file (Markdown) adds simple marks: # for a heading, - for a list item, ** for bold, so people and AI "
    "can see the structure. A skill is a Markdown file called SKILL.md with a label on top: a name and a line that "
    "says when to use it. It sits where the agent looks for skills, so the agent can find it and open it by itself.")

# 5 - The sample prompt -------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
ground(s, gradient=True)
title(s, [("A real prompt that builds", "white")], [("a ", "white"), ("Seer trading bot", "mint")])
CY, CH = 2.1, 4.4
box(s, LM, CY, 7.85, CH, fill="light")
chip = box(s, LM + 0.3, CY + 0.25, 2.25, 0.36, fill="deep", radius=0.18)
label_in(chip, [R("Sample Prompt.docx", 13, font=UISB)])
text(s, LM + 0.3, CY + 0.78, 7.25, 0.62, [R("\u201cI want to create a trading bot to trade Seer prediction markets (seer.pm).\u201d",
                                            16, "deep", font=UISB, italic=True)], label="s5 quote")
rows = [("Read first", [("Seer docs and the Ethereum skills (ethskills.com)", "ink", False)]),
        ("My setup", [("Wallet, chain, bankroll, markets to trade", "ink", False)]),
        ("Build 3 things", [("A skill (SKILL.md)", "purple", True), (", the code, a schedule", "ink", False)]),
        ("Set the rules", [("5-point edge, \u00bc Kelly, max 20% per market", "ink", False)]),
        ("Guard the key", [("Key only in .env, and the AI never reads it", "ink", False)]),
        ("Test first", [("One dry run before any real trade", "ink", False)])]
for j, (lab, runs) in enumerate(rows):
    ry = CY + 1.55 + j * 0.44
    num = box(s, LM + 0.3, ry + 0.02, 0.3, 0.3, fill="deep", shape=MSO_SHAPE.OVAL)
    label_in(num, [R(str(j + 1), 12, bold=True)])
    text(s, LM + 0.78, ry, 1.75, 0.36, [R(lab, 14, "deep", bold=True)], label=f"s5 label {j}")
    text(s, LM + 2.55, ry, 5.1, 0.36, [P(*[R(t, 14, c, bold=b) for t, c, b in runs])], label=f"s5 row {j}")
RX, RW = 8.75, 3.98
box(s, RX, CY, RW, CH, fill="card", line="cardline")
star(s, "mint", RX + 0.3, CY + 0.3, 0.3)
text(s, RX + 0.72, CY + 0.3, RW - 1.0, 0.32, [R("WHY IT WORKS", 14, "mint", bold=True, spc=1.5)], label="s5 why")
text(s, RX + 0.3, CY + 0.85, RW - 0.6, 1.2, [R("It asks the AI to write a skill first.", 24, bold=True)], label="s5 big")
text(s, RX + 0.3, CY + 2.25, RW - 0.6, 2.0, [
    P(R("\u2192 ", 17, "mint", bold=True), R("The method lives in SKILL.md; the code just runs it", 17)),
    P(R("\u2192 ", 17, "mint", bold=True), R("Hard limits that the code enforces", 17), before=12),
    P(R("\u2192 ", 17, "mint", bold=True), R("A dry run before real money", 17), before=12)], label="s5 bullets")
footer(s, 5)
prompt = "\n".join(p.text for p in docx.Document(str(PROMPT_DOCX)).paragraphs)
s.notes_slide.notes_text_frame.text = "Full prompt, from Desktop/Sample Prompt.docx:\n\n" + prompt

prs.save(OUT)
print("wrote", OUT)
